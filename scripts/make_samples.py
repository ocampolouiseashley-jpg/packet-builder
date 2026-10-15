"""Generate the fake sample inputs in samples/.

    python scripts/make_samples.py

Creates a Word template, a CSV of made-up customers, and one PDF "statement
of account" per customer. Two rows are broken on purpose: one is missing a
value and one has no attachment, so the demo shows what gets skipped.
"""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.shared import Pt
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent.parent / "samples"

RECORDS = [
    {"account_id": "ACC-1001", "customer_name": "Maple Street Bakery LLC", "contact_name": "Dana Reyes",
     "address": "14 Maple Street, Springfield, IL 62701", "balance": "1,240.00", "due_date": "2026-11-15"},
    {"account_id": "ACC-1002", "customer_name": "Riverbend Auto Repair", "contact_name": "Sam Okafor",
     "address": "220 River Road, Riverton, WY 82501", "balance": "3,875.50", "due_date": "2026-11-15"},
    {"account_id": "ACC-1003", "customer_name": "Bright Smile Dental", "contact_name": "",  # missing contact
     "address": "9 Clinic Way, Fairview, TN 37062", "balance": "640.00", "due_date": "2026-11-15"},
    {"account_id": "ACC-1004", "customer_name": "Northside Florist", "contact_name": "Priya Nair",
     "address": "77 Bloom Ave, Northfield, MN 55057", "balance": "512.25", "due_date": "2026-11-15"},
]
NO_ATTACHMENT = {"ACC-1004"}


def make_template(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.sections[0].header.paragraphs[0].text = "Example Supply Co.  |  Accounts Receivable  |  Ref: {{account_id}}"

    doc.add_paragraph("{{customer_name}}\nAttn: {{contact_name}}\n{{address}}")
    doc.add_paragraph("Re: Past-due balance on account {{account_id}}")
    doc.add_paragraph("Dear {{contact_name}},")

    # Built from several runs on purpose: Word often stores one placeholder
    # split across runs, and the filler has to handle that.
    p = doc.add_paragraph("Our records show a balance of $")
    p.add_run("{{bal")
    p.add_run("ance}}").bold = True
    p.add_run(" on your account. Please pay by {{due_date}}. A statement of account is enclosed.")

    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Account", "Balance", "Due")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("{{account_id}}", "${{balance}}", "{{due_date}}")):
        cell.text = text

    doc.add_paragraph("If you have already paid, please disregard this letter.")
    doc.add_paragraph("Sincerely,\nAccounts Receivable Team\nExample Supply Co.")
    doc.save(str(path))


def make_statement(record: dict, path: Path) -> None:
    pdf = FPDF(format="Letter")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Statement of Account", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    for line in (
        f"Customer: {record['customer_name']}",
        f"Account: {record['account_id']}",
        "Invoice INV-501   2026-08-01   $" + record["balance"],
        f"Balance due: ${record['balance']}",
    ):
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def main() -> None:
    (ROOT / "attachments").mkdir(parents=True, exist_ok=True)
    make_template(ROOT / "letter_template.docx")
    with (ROOT / "customers.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(RECORDS[0]))
        writer.writeheader()
        writer.writerows(RECORDS)
    for record in RECORDS:
        if record["account_id"] not in NO_ATTACHMENT:
            make_statement(record, ROOT / "attachments" / f"{record['account_id']}.pdf")
    print(f"Samples written to {ROOT}")


if __name__ == "__main__":
    main()
