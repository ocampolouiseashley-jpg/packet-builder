from docx import Document

from packet_builder.fill import all_paragraphs, fill_template, find_placeholders

VALUES = {"ref": "A-1", "first_name": "Dana", "due": "Nov 15"}


def build_template(path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Ref {{ref}}"
    split = doc.add_paragraph("Hello ")
    split.add_run("{{fir")
    split.add_run("st_name}}").italic = True
    split.add_run(", you owe ")
    split.add_run("$100").bold = True
    split.add_run(" by {{due}}.")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Account {{ref}}"
    doc.add_paragraph("Untouched {{not_in_data}}")
    doc.save(str(path))
    return path


def text_of(path):
    return [p.text for p in all_paragraphs(Document(str(path))) if p.text]


def test_find_placeholders_everywhere(tmp_path):
    template = build_template(tmp_path / "t.docx")
    assert find_placeholders(template) == {"ref", "first_name", "due", "not_in_data"}


def test_fills_split_runs_tables_and_headers(tmp_path):
    out = fill_template(build_template(tmp_path / "t.docx"), VALUES, tmp_path / "o.docx")
    text = text_of(out)
    assert "Hello Dana, you owe $100 by Nov 15." in text
    assert "Account A-1" in text
    assert "Ref A-1" in text


def test_other_formatting_is_kept(tmp_path):
    out = fill_template(build_template(tmp_path / "t.docx"), VALUES, tmp_path / "o.docx")
    paragraph = next(p for p in Document(str(out)).paragraphs if p.text.startswith("Hello"))
    assert [r.text for r in paragraph.runs if r.bold] == ["$100"]


def test_unknown_placeholder_is_left_visible(tmp_path):
    out = fill_template(build_template(tmp_path / "t.docx"), {"ref": "A-1"}, tmp_path / "o.docx")
    assert "Untouched {{not_in_data}}" in text_of(out)


def test_two_placeholders_in_one_run(tmp_path):
    doc = Document()
    doc.add_paragraph("{{a}} and {{b}}")
    doc.save(str(tmp_path / "t.docx"))
    out = fill_template(tmp_path / "t.docx", {"a": "first", "b": "second"}, tmp_path / "o.docx")
    assert text_of(out) == ["first and second"]
