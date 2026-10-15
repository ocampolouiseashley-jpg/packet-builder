from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF


def make_pdf(path: Path, pages: int, label: str = "page") -> Path:
    pdf = FPDF(format="Letter")
    pdf.set_font("Helvetica", "", 12)
    for i in range(1, pages + 1):
        pdf.add_page()
        pdf.cell(0, 10, f"{label} {i}")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


@pytest.fixture
def pdf_factory(tmp_path):
    return lambda name, pages, label="page": make_pdf(tmp_path / name, pages, label)


def fake_converter(pages: int = 3):
    """Stands in for LibreOffice: writes an N-page PDF labelled with the start of the DOCX text."""

    def convert(docx_path: Path, out_dir: Path) -> Path:
        text = " ".join(p.text for p in Document(str(docx_path)).paragraphs)
        out = out_dir / (docx_path.stem + ".pdf")
        make_pdf(out, pages, label=text[:40].replace("\n", " ") or "letter")
        return out

    return convert
